import os
import ctypes
import ctypes.wintypes
from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
from PIL import Image
import io
import base64
import json
import openpyxl

app = Flask(__name__, static_folder='static')
CORS(app)

ZEBRA_DEVICE_PATH = r"\\?\USB#VID_0A5F&PID_0065#JAY247698#{28d78fad-5a12-11d1-ae5b-0000f803a8c2}"
GENERIC_WRITE = 0x40000000
FILE_SHARE_READ = 0x00000001
FILE_SHARE_WRITE = 0x00000002
OPEN_EXISTING = 3
INVALID_HANDLE_VALUE = -1

class ZebraUSB:
    def __init__(self, device_path: str):
        self.device_path = device_path
        self.handle = None
        self.kernel32 = ctypes.windll.kernel32

    def open(self) -> bool:
        self.handle = self.kernel32.CreateFileW(
            self.device_path, GENERIC_WRITE, FILE_SHARE_READ | FILE_SHARE_WRITE,
            None, OPEN_EXISTING, 0, None
        )
        return self.handle != INVALID_HANDLE_VALUE

    def write(self, data: bytes) -> bool:
        if not self.handle or self.handle == INVALID_HANDLE_VALUE: return False
        bytes_written = ctypes.wintypes.DWORD(0)
        buf = ctypes.create_string_buffer(data)
        return self.kernel32.WriteFile(self.handle, buf, len(data), ctypes.byref(bytes_written), None)

    def close(self):
        if self.handle and self.handle != INVALID_HANDLE_VALUE:
            self.kernel32.CloseHandle(self.handle)
            self.handle = None

    def __enter__(self):
        self.open()
        return self

    def __exit__(self, *args):
        self.close()

def image_to_zpl(base64_str, x, y, width=None, height=None):
    if "," in base64_str:
        base64_str = base64_str.split(",")[1]
    
    try:
        image_data = base64.b64decode(base64_str)
        img = Image.open(io.BytesIO(image_data))
        
        if width and height:
            img = img.resize((int(width), int(height)), Image.Resampling.LANCZOS)
        
        img = img.convert("1")
        
        width_px, height_px = img.size
        bytes_per_row = (width_px + 7) // 8
        total_bytes = bytes_per_row * height_px
        
        hex_data = ""
        for y_pos in range(height_px):
            row_bytes = bytearray(bytes_per_row)
            for x_pos in range(width_px):
                pixel = img.getpixel((x_pos, y_pos))
                if pixel == 0:  # Black pixel
                    byte_idx = x_pos // 8
                    bit_idx = 7 - (x_pos % 8)
                    row_bytes[byte_idx] |= (1 << bit_idx)
            
            hex_data += "".join(f"{b:02X}" for b in row_bytes) + "\n"
            
        return f"^FO{x},{y}^GFA,{total_bytes},{total_bytes},{bytes_per_row},{hex_data}^FS"
    except Exception as e:
        print("Image processing error:", e)
        return ""

def build_zpl(elements, canvas_width=1189, canvas_height=795):
    zpl = ["^XA", "^CI28", f"^PW{canvas_width}", f"^LL{canvas_height}"]
    for el in elements:
        x = int(el.get('x', 0))
        y = int(el.get('y', 0))
        el_type = el.get('type')
        if el_type == 'text':
            size = int(el.get('fontSize', 30))
            text = str(el.get('text', ''))
            zpl.append(f"^FO{x},{y}^A0N,{size},{size}^FD{text}^FS")
        elif el_type == 'barcode':
            data = str(el.get('data', ''))
            height = int(el.get('height', 100))
            zpl.append(f"^FO{x},{y}^BY3,3,{height}^BCN,{height},Y,N,N^FD{data}^FS")
        elif el_type == 'qrcode':
            data = str(el.get('data', ''))
            mag = int(el.get('size', 5))
            zpl.append(f"^FO{x},{y}^BQN,2,{mag}^FDQA,{data}^FS")
        elif el_type == 'box':
            w = int(el.get('width', 100))
            h = int(el.get('height', 100))
            t = int(el.get('thickness', 2))
            zpl.append(f"^FO{x},{y}^GB{w},{h},{t}^FS")
        elif el_type == 'image':
            base64_data = el.get('src', '')
            w = el.get('width')
            h = el.get('height')
            if base64_data:
                zpl_img = image_to_zpl(base64_data, x, y, w, h)
                if zpl_img:
                    zpl.append(zpl_img)
    zpl.append("^PQ1")
    zpl.append("^XZ")
    return "\n".join(zpl)

@app.route('/')
def index():
    return send_from_directory('static', 'index.html')

@app.route('/print', methods=['POST'])
def print_label():
    data = request.json
    if not isinstance(data, dict) or 'elements' not in data:
        return jsonify({"success": False, "error": "Invalid payload format"}), 400
    
    elements = data.get('elements', [])
    canvas_w = int(data.get('canvasWidth', 1189))
    canvas_h = int(data.get('canvasHeight', 795))
    
    zpl = build_zpl(elements, canvas_w, canvas_h)
    print("Generating ZPL:\n", zpl)
    
    try:
        with ZebraUSB(ZEBRA_DEVICE_PATH) as printer:
            if not printer.handle or printer.handle == INVALID_HANDLE_VALUE:
                return jsonify({"success": False, "error": "Could not connect to Zebra printer (Check USB and Power)"}), 500
            
            if printer.write(zpl.encode("utf-8", errors="replace")):
                return jsonify({"success": True})
            else:
                return jsonify({"success": False, "error": "Failed to write to printer"}), 500
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


TEMPLATES_DIR = "templates"
os.makedirs(TEMPLATES_DIR, exist_ok=True)
EXCEL_FILE = "Etiquetas.xlsx"

def clean_excel_val(value) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    if s.startswith("_"):
        s = s[1:]
    return s

def read_equipment_data(filepath: str):
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    records = []

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        vals = [cell.value for cell in row]
        if not vals[0] or not vals[1]:
            continue
        if vals[3] and "NO HAY" in str(vals[3]).upper():
            continue

        records.append({
            "row_num":   row[0].row,
            "ubicacion": clean_excel_val(vals[0]),
            "id_equipo": clean_excel_val(vals[1]),
            "tipo":      clean_excel_val(vals[2]),
            "marca":     clean_excel_val(vals[3]),
            "modelo":    clean_excel_val(vals[4]),
            "serial":    clean_excel_val(vals[5]),
            "usuario":   clean_excel_val(vals[6]),
            "username":  clean_excel_val(vals[7]),
            "area":      clean_excel_val(vals[8]),
            "area_code": clean_excel_val(vals[9]),
            "estado":    clean_excel_val(vals[10]),
        })
    wb.close()
    return records

def generate_asset_zpl(row_data: dict) -> str:
    ubicacion = row_data["ubicacion"]
    id_equipo = row_data["id_equipo"]
    tipo      = row_data["tipo"]
    marca     = row_data["marca"]
    modelo    = row_data["modelo"]
    serial    = row_data["serial"]
    usuario   = row_data["usuario"]
    area      = row_data["area"]
    area_code = row_data["area_code"]
    estado    = row_data["estado"]

    barcode_data = f"{ubicacion}-{area_code}-{id_equipo}"

    tipo_display = {
        "PC": "PC", "LTP": "LAPTOP", "IMP": "IMPRESORA",
        "MON": "MONITOR", "SRV": "SERVER",
    }.get(tipo, tipo)

    zpl = f'''^XA
^CI28
^PW1189
^LL795
^FO40,60^GB1109,0,4^FS
^FO40,75^A0N,50,50^FD ETIQUETA DE ACTIVO^FS
^FO40,140^GB1109,0,2^FS
^FO40,165^A0N,28,28^FDUbicacion:^FS
^FO260,165^A0N,32,32^FD{ubicacion}^FS
^FO600,165^A0N,28,28^FDID:^FS
^FO680,165^A0N,32,32^FD{id_equipo}^FS
^FO40,215^A0N,28,28^FDTipo:^FS
^FO260,215^A0N,32,32^FD{tipo_display}^FS
^FO600,215^A0N,28,28^FDMarca:^FS
^FO700,215^A0N,32,32^FD{marca[:25]}^FS
^FO40,265^A0N,28,28^FDModelo:^FS
^FO260,265^A0N,32,32^FD{modelo[:40]}^FS
^FO40,315^A0N,28,28^FDSerie:^FS
^FO260,315^A0N,32,32^FD{serial[:35]}^FS
^FO40,365^A0N,28,28^FDEstado:^FS
^FO260,365^A0N,32,32^FD{estado}^FS
^FO40,410^GB1109,0,2^FS
^FO40,435^A0N,28,28^FDUsuario:^FS
^FO200,435^A0N,32,32^FD{usuario[:45]}^FS
^FO40,485^A0N,28,28^FDArea:^FS
^FO200,485^A0N,32,32^FD{area} ({area_code})^FS
^FO40,530^GB1109,0,2^FS
^FO40,555^BY2,3,100^BCN,100,Y,N,N^FD{barcode_data}^FS
^FO40,685^A0N,22,22^FD{barcode_data}^FS
^FO40,720^GB1109,0,4^FS
^PQ1
^XZ'''
    return zpl


import werkzeug.utils

@app.route('/api/templates', methods=['GET', 'POST', 'DELETE'])
def manage_templates():
    if request.method == 'GET':
        templates = []
        for f in os.listdir(TEMPLATES_DIR):
            if f.endswith('.json'):
                templates.append(f[:-5])
        return jsonify({"success": True, "templates": templates})

    elif request.method == 'POST':
        data = request.json
        name = data.get('name')
        if not name:
            return jsonify({"success": False, "error": "Template name required"}), 400

        safe_name = werkzeug.utils.secure_filename(name)
        if not safe_name:
             return jsonify({"success": False, "error": "Invalid template name"}), 400

        filepath = os.path.join(TEMPLATES_DIR, f"{safe_name}.json")
        with open(filepath, 'w') as f:
            json.dump(data.get('elements', []), f)
        return jsonify({"success": True})

    elif request.method == 'DELETE':
        name = request.args.get('name')
        if not name:
            return jsonify({"success": False, "error": "Template name required"}), 400

        safe_name = werkzeug.utils.secure_filename(name)
        if not safe_name:
             return jsonify({"success": False, "error": "Invalid template name"}), 400

        filepath = os.path.join(TEMPLATES_DIR, f"{safe_name}.json")
        if os.path.exists(filepath):
            os.remove(filepath)
            return jsonify({"success": True})
        return jsonify({"success": False, "error": "Template not found"}), 404

@app.route('/api/templates/<name>', methods=['GET'])
def get_template(name):
    safe_name = werkzeug.utils.secure_filename(name)
    if not safe_name:
         return jsonify({"success": False, "error": "Invalid template name"}), 400

    filepath = os.path.join(TEMPLATES_DIR, f"{safe_name}.json")
    if not os.path.exists(filepath):
        return jsonify({"success": False, "error": "Template not found"}), 404

    with open(filepath, 'r') as f:
        elements = json.load(f)
    return jsonify({"success": True, "elements": elements})


@app.route('/api/equipment', methods=['GET'])
def get_equipment():
    if not os.path.exists(EXCEL_FILE):
        return jsonify({"success": False, "error": "Excel file not found"}), 404

    try:
        records = read_equipment_data(EXCEL_FILE)
        return jsonify({"success": True, "data": records})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/print_asset', methods=['POST'])
def print_asset():
    data = request.json
    row_nums = data.get('rows', [])
    if not row_nums:
        return jsonify({"success": False, "error": "No rows specified"}), 400

    if not os.path.exists(EXCEL_FILE):
        return jsonify({"success": False, "error": "Excel file not found"}), 404

    try:
        records = read_equipment_data(EXCEL_FILE)
        to_print = [r for r in records if r["row_num"] in row_nums]

        if not to_print:
            return jsonify({"success": False, "error": "No matching records found"}), 404

        with ZebraUSB(ZEBRA_DEVICE_PATH) as printer:
            if not printer.handle or printer.handle == INVALID_HANDLE_VALUE:
                return jsonify({"success": False, "error": "Could not connect to Zebra printer"}), 500

            success_count = 0
            for r in to_print:
                zpl = generate_asset_zpl(r)
                if printer.write(zpl.encode("utf-8", errors="replace")):
                    success_count += 1

            if success_count == len(to_print):
                return jsonify({"success": True, "message": f"Printed {success_count} labels"})
            else:
                return jsonify({"success": False, "error": f"Failed to print some labels. {success_count}/{len(to_print)} printed."}), 500

    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/zpl_preview', methods=['POST'])
def zpl_preview():
    data = request.json
    if not isinstance(data, dict) or 'elements' not in data:
        return jsonify({"success": False, "error": "Invalid payload format"}), 400

    elements = data.get('elements', [])
    canvas_w = int(data.get('canvasWidth', 1189))
    canvas_h = int(data.get('canvasHeight', 795))

    zpl = build_zpl(elements, canvas_w, canvas_h)
    return jsonify({"success": True, "zpl": zpl})

if __name__ == '__main__':

    print("Starting Zebra Label Editor Server...")
    print("Open http://localhost:5000 in your browser.")
    app.run(host='0.0.0.0', port=5000)
